"""Generate concise and detailed STM-GARSOS presentations.

Outputs:
  STM_GARSOS_concise.tex
  STM_GARSOS_detailed.tex
  STM_GARSOS_concise.pptx
  STM_GARSOS_detailed.pptx
"""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parent

CONCISE = [
    ("STM-GARSOS processing", ["From GAR-SOS radial data to STM spatial coefficients",
                               "Concise overview of the historical stable workflow"]),
    ("Objective and model", [
        r"Represent each dynamic image as $X_t(x,y)=\sum_{\ell=1}^{L} {\rm ST}_{t,\ell}(x,y) C_\ell(x,y)$.",
        "Estimate L spatial coefficient images from multicoil radial data.",
        "Synthesize the complete dynamic image series after the coefficient solve.",
    ]),
    ("Raw-data preparation", [
        "Load GAR-SOS metadata, raw k-space, and a vendor-reconstructed reference volume.",
        "Correct readout delay and prepare partition data.",
        "Estimate/apply coil-noise whitening.",
        "Generate normalized [kx, ky, kz] trajectory coordinates.",
        "Compute radial density-compensation weights.",
    ]),
    ("Frame formation", [
        "Flattened arrays retain one row per readout/partition/spoke sample.",
        "Spokes are grouped consecutively into temporal frames.",
        "Historical baseline: 15 spokes per frame and approximately 133 complete frames.",
        "Incomplete trailing spokes are discarded.",
    ]),
    ("Calibration-data path", [
        "Each frame is reconstructed with 3-D Kaiser-Bessel gridding.",
        "Oversampled grid: approximately 308 x 308 x 132.",
        "3-D inverse FFT and centered resize produce a 224 x 224 x 96 volume.",
        "A selected z slice is transformed to 2-D Cartesian k-space.",
        "A central 32 x 32 calibration square is extracted.",
    ]),
    ("STM calibration", [
        r"$k_{\rm Cal}$ is supplied to $\texttt{STM\_computation}$.",
        "The calibration data span the temporal frames.",
        "The STM algorithm estimates ST maps and eigenvalues.",
        "Historical baseline: L = 7 and 224 x 224 spatial map dimensions.",
    ]),
    ("Radial STM data", [
        "The historical baseline used a direct centered 58-point partition IFFT.",
        "One z index is selected and the coil dimension is preserved.",
        "Output radial data: [readout x spokes/frame x coils x frames].",
        "In-plane kx/ky coordinates are retained for the NUFFT.",
    ]),
    ("Multicoil radial operator", [
        r"Forward: $d_c={\cal N}_{\rm radial}(S_c \odot X)$.",
        r"Adjoint: ${\cal F}^H(d)=\sum_c S_c^* \odot {\cal N}_{\rm radial}^H(d_c)$.",
        "Sensitivity maps are forward receive-coil maps.",
        "Coils remain separate in the data fidelity term.",
    ]),
    ("STM operator and PCG", [
        "STM synthesis maps coefficients C to dynamic images X.",
        "The radial operator maps X to predicted multicoil radial data.",
        r"Solve $({\cal A}^H{\cal A})C={\cal A}^H y$ with PCG.",
        "Optional Tikhonov solve adds lambda I.",
    ]),
    ("Output and historical caveats", [
        "Output: coefficient images plus synthesized [224 x 224 x frames] image series.",
        "The initial baseline used the direct 58-point z/partition IFFT.",
        "Later work addresses forward-map estimation, z-grid consistency, scaling, and conditioning.",
        "The slides describe the original stable path, not all later corrections.",
    ]),
]

DETAILED = [
    ("STM-GARSOS processing", ["Detailed methods presentation",
                               "Historical stable workflow through spatial-coefficient estimation"]),
    ("Reconstruction objective", [
        r"Dynamic image model: $X_t(x,y)=\sum_{\ell=1}^{L}{\rm ST}_{t,\ell}(x,y)C_\ell(x,y)$.",
        "Unknown: C_l(x,y), an [Nx x Ny x L] coefficient volume.",
        "Measured data: multicoil radial samples for each temporal frame.",
        "Goal: fit the measured radial data while restricting temporal variation to the STM subspace.",
    ]),
    ("Entry point and inputs", [
        "main_STM_GARSOS loads input k-space identifiers and scan metadata.",
        "A vendor-reconstructed reference volume defines target geometry.",
        "prepareGARSOSSTMData creates shared prepared data and calibration metadata.",
        "Historical defaults: 15 spokes/frame, 32 x 32 calibration square, L = 7.",
    ]),
    ("Prepared trajectory and data", [
        "sampledData: [nReadout x nPartitions x nSpokes, nCoils].",
        "kSpaceLocations: flattened [kx, ky, kz] rows.",
        "spokeIndex: one identifier per flattened sample.",
        "DCF: one radial density-compensation value per sample.",
        "Coordinates use normalized cycles/FOV-like units.",
    ]),
    ("Trajectory construction", [
        r"$k_x=\rho\cos(\phi)$ and $k_y=\rho\sin(\phi)$.",
        "rho is the centered readout coordinate normalized by readout length.",
        "phi follows the GAR-SOS spoke-angle convention.",
        "kz is generated from the partition index and acquisition geometry.",
        "The arrays are flattened in readout/partition/spoke order.",
    ]),
    ("Readout processing and whitening", [
        "Readout delay correction is estimated from selected calibration lines.",
        "Readout Fourier processing is applied before the final flattening.",
        "Partition padding is performed during preparation.",
        "Noise covariance is estimated from noise acquisitions.",
        "The whitening transform is applied across receive channels.",
    ]),
    ("Density compensation", [
        "The existing DCF is based on radial shell-sector area.",
        "It is generated from the in-plane radius and readout increment.",
        "Frame-specific DCF values are selected with the frame samples.",
        "Calibration gridding applies DCF and additional reconstruction weights.",
        "The later STM radial operator uses its prepared radial samples and NUFFT convention.",
    ]),
    ("Frame grouping", [
        "Unique spoke identifiers define the acquisition order.",
        "nFrames = floor(nSpokes / nSpokesPerFrame).",
        "Each frame contains all readout and partition samples for its spokes.",
        "Incomplete trailing spokes are excluded.",
        "For the baseline, 15 spokes/frame produced approximately 133 frames.",
    ]),
    ("Calibration reconstruction", [
        "reconstructGARFrameForSTM selects frameData, frameTrajectory, and frameDCF.",
        "Coordinate transform and translation phase correction are applied.",
        "Gaussian/spatial weighting and DCF scaling are applied.",
        "GpuReconstructur20181008 grids every coil separately.",
        "The class applies inverse FFT, resize/deapodization, and coil-profile accumulation.",
    ]),
    ("3-D grid and target geometry", [
        "Target image geometry: imageSize = [224, 224, 96].",
        "Oversampled grid: gridSize = ceil(1.375 x imageSize) = [308, 308, 132].",
        "Samples are deposited with a 3-D Kaiser-Bessel kernel.",
        "ifftn is applied on the oversampled Cartesian grid.",
        "shiftAndResize maps the result to the centered 224 x 224 x 96 grid.",
    ]),
    ("Calibration k-space extraction", [
        "Select reconstructedVolume(:, :, zSlice).",
        "Apply fftshift(fft2(ifftshift(senseImage))).",
        "Extract the centered calibration square with localCentralRegion.",
        "Baseline calibrationSize = [32, 32].",
        "Single-slice kCal has dimensions [32 x 32 x nFrames].",
    ]),
    ("STM computation", [
        r"$\texttt{STM\_computation}$ receives $k_{\rm Cal}$, desired map dimensions, $L$, and options.",
        "The calibration data encode temporal correlations across frames.",
        "The algorithm estimates ST_maps [Nx x Ny x Nt x L].",
        "eigenValues provide a diagnostic of model rank/content.",
        "The baseline used L = 7 because more modes appeared informative than in the Cartesian demo.",
    ]),
    ("Historical z-slice preparation", [
        "prepareGARSOSSliceDataForMCNUFFT restores [readout, partition, spoke, coil].",
        "It applies the selected translation phase and coordinate transform.",
        "Historical baseline: centered direct IFFT along the 58 partition entries.",
        "The requested zIndex is selected from that 58-point result.",
        "This differs from the later 58-to-96 gridded z-only correction.",
    ]),
    ("Radial STM data dimensions", [
        "After frame grouping: radialData [nReadout x nSpokesPerFrame x nCoil x nFrame].",
        "kxy [nReadout x nSpokesPerFrame x nFrame].",
        "The in-plane radial trajectory is retained.",
        "The coil dimension is not combined before the STM solve.",
        "Trajectory/data correspondence is validated before operator construction.",
    ]),
    ("Forward sensitivity-map model", [
        r"Forward coil model: $d_c={\cal N}_{\rm radial}(S_c\odot X)$.",
        "S_c are forward receive sensitivities in the whitened channel basis.",
        r"Adjoint coil combination is $\sum_c S_c^* d_c$ after NUFFT adjoints.",
        "This is distinct from calibration-only coil-combination profiles.",
        "Map and dynamic-data channel ordering must match.",
    ]),
    ("Radial NUFFT operator", [
        "buildGARSOSRadialEncodingOperator creates one NUFFT plan per frame.",
        "Forward: multiply by coil maps, then NUFFT each coil.",
        "Adjoint: NUFFT-adjoint each coil, then sum conjugated maps.",
        "Optional single precision converts trajectories, maps, data, and outputs.",
        "The operator exposes forward, adjoint, and vectorized forms.",
    ]),
    ("STM operator composition", [
        "buildGARSOSSTMOperators composes STM synthesis with radial encoding.",
        r"$\mathcal{A}(C)=\mathcal{F}(\mathcal{S}(C))$.",
        r"$\mathcal{A}^H(y)=\mathcal{S}^H(\mathcal{F}^H(y))$.",
        "The normal operator is applied without explicitly forming a matrix.",
        "The coefficient vector has Nx x Ny x L entries.",
    ]),
    ("PCG solve", [
        r"Right-hand side: $b=\mathcal{A}^H y$.",
        r"Unregularized system: $(\mathcal{A}^H\mathcal{A})C=b$.",
        "The baseline used tolerance 1e-5 and up to 50 iterations.",
        "Residual history, status flag, iteration count, and elapsed time are stored.",
        "Preconditioning and Toeplitz support were preliminary rather than baseline assumptions.",
    ]),
    ("Optional Tikhonov solve", [
        r"Regularized system: $(\mathcal{A}^H\mathcal{A}+\lambda I)C=b$.",
        "The implementation can run a second PCG solve with lambda.",
        "The effective lambda depends on data, map, NUFFT, and DCF scaling.",
        "A numerical lambda should not be compared across differently scaled operators without calibration.",
    ]),
    ("Synthesis and outputs", [
        r"$\widehat{X}_t=\sum_{\ell=1}^{L}{\rm ST}_{t,\ell}\widehat{C}_\ell$.",
        "sense_recon_stm has dimensions [Nx x Ny x nFrame].",
        "The result also contains coefficient images and PCG diagnostics.",
        "The initial reported output contained approximately 133 reconstructed frames.",
    ]),
    ("Baseline versus later improvements", [
        "Baseline: direct 58-point partition IFFT for STM slice data.",
        "Later work: forward sensitivity-map estimator and map-semantics audit.",
        "Later work: coil-preserving gridded z-only path for the 96-slice geometry.",
        "Later work: scaling, conditioning, coverage, regularization, and preconditioning analysis.",
        "These improvements should be presented separately from the historical baseline.",
    ]),
    ("End-to-end summary", [
        "Raw GAR-SOS data -> corrected, whitened radial samples.",
        "Trajectory + DCF + spoke IDs -> temporal frames.",
        "3-D gridding -> calibration Cartesian k-space -> central square.",
        "STM calibration -> ST maps and eigenvalues.",
        "Coil-resolved radial data + sensitivity maps -> STM/radial operator.",
        "PCG -> spatial coefficients -> dynamic image series.",
    ]),
]

def latex_escape(text):
    replacements = {"&": r"\&", "%": r"\%", "_": r"\_", "#": r"\#"}
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text

def write_tex(slides, path, title):
    lines = [
        r"\documentclass[aspectratio=169,professionalfonts]{beamer}",
        r"\usetheme{Madrid}",
        r"\usecolortheme{default}",
        r"\usepackage{amsmath,amssymb}",
        r"\usepackage[T1]{fontenc}",
        r"\title{" + title + r"}",
        r"\author{STM-GARSOS workflow}",
        r"\date{\today}",
        r"\begin{document}",
        r"\begin{frame}\titlepage\end{frame}",
    ]
    for title_text, bullets in slides[1:]:
        lines += [r"\begin{frame}{" + title_text + r"}", r"\begin{itemize}"]
        for bullet in bullets:
            if "$" not in bullet:
                bullet = bullet.replace("_", r"\_")
            lines.append(r"\item " + bullet)
        lines += [r"\end{itemize}", r"\end{frame}"]
    lines += [r"\end{document}"]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")

def add_ppt_slide(prs, title, bullets, title_slide=False):
    layout = prs.slide_layouts[0 if title_slide else 1]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title
    if title_slide:
        slide.placeholders[1].text = "\n".join(bullets)
    else:
        body = slide.placeholders[1].text_frame
        body.clear()
        for i, bullet in enumerate(bullets):
            paragraph = body.paragraphs[0] if i == 0 else body.add_paragraph()
            paragraph.text = bullet.replace("$", "")
            paragraph.level = 0
            paragraph.font.size = Pt(21 if len(bullets) <= 4 else 18)
    return slide

def write_ppt(slides, path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for i, (title, bullets) in enumerate(slides):
        slide = add_ppt_slide(prs, title, bullets, title_slide=(i == 0))
        for shape in slide.shapes:
            if hasattr(shape, "text_frame"):
                for paragraph in shape.text_frame.paragraphs:
                    paragraph.font.name = "Aptos"
        if i > 0:
            bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(.18), prs.slide_height)
            bar.fill.solid()
            bar.fill.fore_color.rgb = RGBColor(31, 78, 121)
            bar.line.fill.background()
    prs.save(path)

if __name__ == "__main__":
    write_tex(CONCISE, ROOT / "STM_GARSOS_concise.tex", "STM-GARSOS Processing: Concise Overview")
    write_tex(DETAILED, ROOT / "STM_GARSOS_detailed.tex", "STM-GARSOS Processing: Detailed Methods")
    write_ppt(CONCISE, ROOT / "STM_GARSOS_concise.pptx")
    write_ppt(DETAILED, ROOT / "STM_GARSOS_detailed.pptx")
    print("Generated four presentation files in", ROOT)
